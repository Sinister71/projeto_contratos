"""
Resumo semanal de contratos nas faixas de vencimento (30 / 60 / 90 dias).

Sem dependência do Streamlit: pode ser usado pela tela do app e também pelo
script agendável `gerar_resumo_semanal.py`.

"Entraram na faixa nesta semana" = contratos que cruzaram o limite da faixa nos
últimos 7 dias. Como os dias restantes diminuem 1 por dia, isso equivale a estar
nos 7 últimos dias da faixa (ex.: de 24 a 30 dias restantes para "Até 30 dias").
Rodando o resumo sempre no mesmo dia da semana, cada contrato aparece como
"novidade" exatamente uma vez por faixa.
"""

import html
import io
import os
from datetime import date

import pandas as pd  # type: ignore
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore
from openpyxl.utils import get_column_letter  # type: ignore

from database import PASTA_PROJETO
from queries_contratos import classificar_faixa, listar_contratos

JANELA_DIAS = 7

# Constantes para evitar duplicação de literais (S1192)
FAIXA_ATE_30 = "Até 30 dias"
FAIXA_31_60 = "31 a 60 dias"
FAIXA_61_90 = "61 a 90 dias"
FAIXA_VENCIDO = "Vencido"
COLUNA_TOTAL_ALERTA = "Total em alerta"

# (nome da faixa, limite superior em dias)
FAIXAS_ALERTA = [(FAIXA_ATE_30, 30), (FAIXA_31_60, 60), (FAIXA_61_90, 90)]
FAIXAS_EXIBIDAS = [FAIXA_VENCIDO] + [f for f, _ in FAIXAS_ALERTA]

ICONES = {FAIXA_VENCIDO: "🔴", FAIXA_ATE_30: "🟠", FAIXA_31_60: "🟡", FAIXA_61_90: "🟢"}
COR_FAIXA = {FAIXA_VENCIDO: "A32D2D", FAIXA_ATE_30: "E24B4A", FAIXA_31_60: "EF9F27", FAIXA_61_90: "5DAA6E"}

COLUNAS = ["Tipo", "Nº", "Fornecedor", "Objeto", "Vencimento", "Dias", "Fiscal"]
COLUNAS_CURTAS = {"Tipo", "Nº", "Vencimento", "Dias", "Faixa"}
LARGURA_COLUNA = {"Tipo": 70, "Nº": 80, "Vencimento": 90, "Dias": 50, "Faixa": 100}


# ---------------------------------------------------------------------------
# Dados
# ---------------------------------------------------------------------------
def _tabela(sub: pd.DataFrame) -> pd.DataFrame:
    sub = sub.sort_values("dias_para_vencer", na_position="last")
    return pd.DataFrame(
        {
            "Tipo": sub["tipo"].replace({"Ata de registro de preços": "ARP"}),
            "Nº": sub["numero_contrato"],
            "Fornecedor": sub["razao_social"],
            "Objeto": sub["objeto"],
            "Vencimento": sub["data_fim"].dt.strftime("%d/%m/%Y").fillna("—"),
            "Dias": sub["dias_para_vencer"].astype("Int64"),
            "Fiscal": sub["fiscal"].replace("", pd.NA).fillna("—"),
        },
        columns=COLUNAS,
    ).reset_index(drop=True)


def montar_resumo(df: pd.DataFrame | None = None, hoje=None) -> dict:
    """Calcula tudo o que o resumo semanal mostra. `hoje` permite simular outras datas."""
    hoje = pd.Timestamp(hoje or date.today()).normalize()
    df = (listar_contratos() if df is None else df).copy()

    if df.empty:
        return {
            "hoje": hoje,
            "total": 0,
            "contagens": dict.fromkeys(FAIXAS_EXIBIDAS, 0),
            "faixas": {f: pd.DataFrame(columns=COLUNAS) for f in FAIXAS_EXIBIDAS},
            "entraram": {f: pd.DataFrame(columns=COLUNAS) for f in FAIXAS_EXIBIDAS},
            "vencem_semana": pd.DataFrame(columns=COLUNAS),
            "por_fiscal": pd.DataFrame(),
        }

    # recalcula em relação a `hoje` (assim o resumo pode ser gerado para qualquer data)
    df["dias_para_vencer"] = (df["data_fim"] - hoje).dt.days
    df["faixa_vencimento"] = df["dias_para_vencer"].map(classificar_faixa)
    dias = df["dias_para_vencer"]

    faixas = {f: _tabela(df[df["faixa_vencimento"] == f]) for f in FAIXAS_EXIBIDAS}

    entraram = {}
    for nome, limite in FAIXAS_ALERTA:
        mascara = (dias >= limite - (JANELA_DIAS - 1)) & (dias <= limite)
        entraram[nome] = _tabela(df[mascara])
    # "Vencido": os que venceram nos últimos 7 dias
    entraram[FAIXA_VENCIDO] = _tabela(df[(dias < 0) & (dias >= -JANELA_DIAS)])

    vencem_semana = _tabela(df[(dias >= 0) & (dias <= JANELA_DIAS)])

    # por fiscal
    base = df[df["faixa_vencimento"].isin(FAIXAS_EXIBIDAS)].copy()
    base["fiscal"] = base["fiscal"].replace("", pd.NA).fillna("(sem fiscal)")
    if base.empty:
        por_fiscal = pd.DataFrame()
    else:
        por_fiscal = (
            base.pivot_table(index="fiscal", columns="faixa_vencimento", values="id",
                             aggfunc="count", fill_value=0)
            .reindex(columns=FAIXAS_EXIBIDAS, fill_value=0)
        )
        por_fiscal[COLUNA_TOTAL_ALERTA] = por_fiscal.sum(axis=1)
        por_fiscal = por_fiscal.sort_values(COLUNA_TOTAL_ALERTA, ascending=False)
        por_fiscal.index.name = "Fiscal"
        por_fiscal.columns.name = None
        por_fiscal = por_fiscal.reset_index()

    return {
        "hoje": hoje,
        "total": len(df),
        "contagens": {f: len(faixas[f]) for f in FAIXAS_EXIBIDAS},
        "faixas": faixas,
        "entraram": entraram,
        "vencem_semana": vencem_semana,
        "por_fiscal": por_fiscal,
    }


# ---------------------------------------------------------------------------
# Texto (Teams / WhatsApp / corpo de e-mail)
# ---------------------------------------------------------------------------
def _prazo(dias) -> str:
    if pd.isna(dias):
        return "sem data"
    dias = int(dias)
    if dias < 0:
        return f"venceu há {-dias} dia" + ("" if dias == -1 else "s")
    if dias == 0:
        return "vence hoje"
    return f"vence em {dias} dia" + ("" if dias == 1 else "s")


def _linha(r) -> str:
    fiscal = "" if r["Fiscal"] == "—" else f" | Fiscal: {r['Fiscal']}"
    return f"  • {r['Tipo']} {r['Nº']} — {r['Fornecedor']} — {r['Vencimento']} ({_prazo(r['Dias'])}){fiscal}"


def gerar_texto(resumo: dict) -> str:
    hoje = resumo["hoje"]
    c = resumo["contagens"]
    linhas = [
        f"📋 RESUMO SEMANAL DE CONTRATOS — {hoje:%d/%m/%Y}",
        "",
        f"{resumo['total']} contratos e atas ativos.",
        "",
    ]
    for f in FAIXAS_EXIBIDAS:
        linhas.append(f"{ICONES[f]} {f}: {c[f]}")

    novidades = [(f, resumo["entraram"][f]) for f in FAIXAS_EXIBIDAS if not resumo["entraram"][f].empty]
    linhas += ["", "🔔 ENTRARAM NA FAIXA NOS ÚLTIMOS 7 DIAS"]
    if not novidades:
        linhas.append("  Nenhum contrato mudou de faixa esta semana.")
    for f, t in novidades:
        titulo = "Venceram" if f == FAIXA_VENCIDO else f"Passaram para {f}"
        linhas.append(f"{ICONES[f]} {titulo}:")
        linhas += [_linha(r) for _, r in t.iterrows()]

    if not resumo["vencem_semana"].empty:
        linhas += ["", "⚠️ VENCEM NOS PRÓXIMOS 7 DIAS"]
        linhas += [_linha(r) for _, r in resumo["vencem_semana"].iterrows()]

    for f in FAIXAS_EXIBIDAS:
        t = resumo["faixas"][f]
        if t.empty:
            continue
        linhas += ["", f"{ICONES[f]} {f.upper()} ({len(t)})"]
        linhas += [_linha(r) for _, r in t.iterrows()]

    pf = resumo["por_fiscal"]
    if not pf.empty:
        linhas += ["", "👤 POR FISCAL (contratos em alerta)"]
        for _, r in pf.iterrows():
            partes = [f"{f.lower()}: {int(r[f])}" for f in FAIXAS_EXIBIDAS if int(r[f])]
            linhas.append(f"  • {r['Fiscal']}: {int(r[COLUNA_TOTAL_ALERTA])} ({', '.join(partes)})")

    return "\n".join(linhas)


# ---------------------------------------------------------------------------
# HTML (abrir no navegador, imprimir em PDF ou colar no Outlook)
# ---------------------------------------------------------------------------
_CSS = """
body{font-family:Arial,Helvetica,sans-serif;color:#222;max-width:1000px;margin:24px auto;padding:0 16px;font-size:14px}
h1{font-size:22px;margin-bottom:2px} h2{font-size:16px;margin:26px 0 8px;border-bottom:1px solid #ddd;padding-bottom:4px}
.sub{color:#666;margin-bottom:16px}
table.kpis{border-collapse:separate;border-spacing:8px 0;width:100%;margin:14px -8px}
table.kpis td{color:#fff;text-align:center;padding:10px 6px;width:25%}
table.kpis b{display:block;font-size:26px}
table.t{border-collapse:collapse;width:100%;margin:6px 0 12px;font-size:13px}
table.t th{color:#fff;text-align:left;padding:6px 8px}
table.t td{border-bottom:1px solid #e3e3e3;padding:5px 8px;vertical-align:top}
table.t tr:nth-child(even) td{background:#f7f7f7}
.vazio{color:#777;font-style:italic}
@media print{body{margin:0}}
"""


def _html_tabela(t: pd.DataFrame, cor: str) -> str:
    if t.empty:
        return '<p class="vazio">Nenhum contrato.</p>'
    cab = ""
    for c in t.columns:
        largura = f' width="{LARGURA_COLUNA[c]}"' if c in LARGURA_COLUNA else ""
        cab += f'<th bgcolor="#{cor}"{largura} style="background:#{cor}">{html.escape(c)}</th>'

    corpo = ""
    for _, r in t.iterrows():
        celulas = ""
        for coluna, v in r.items():
            texto = html.escape("" if pd.isna(v) else str(v))
            estilo = ' style="white-space:nowrap"' if coluna in COLUNAS_CURTAS else ""
            celulas += f"<td{estilo}>{texto}</td>"
        corpo += f"<tr>{celulas}</tr>"
    return f'<table class="t" width="100%"><tr>{cab}</tr>{corpo}</table>'


def gerar_html(resumo: dict) -> str:
    hoje = resumo["hoje"]
    c = resumo["contagens"]
    kpis = "".join(
        f'<td bgcolor="#{COR_FAIXA[f]}" style="background:#{COR_FAIXA[f]}"><b>{c[f]}</b>{f}</td>'
        for f in FAIXAS_EXIBIDAS
    )
    partes = [
        "<h1>Resumo semanal de contratos</h1>",
        f'<div class="sub">Posição em {hoje:%d/%m/%Y} — {resumo["total"]} contratos e atas ativos</div>',
        f'<table class="kpis"><tr>{kpis}</tr></table>',
        "<h2>🔔 Entraram na faixa nos últimos 7 dias</h2>",
    ]
    algum = False
    for f in FAIXAS_EXIBIDAS:
        t = resumo["entraram"][f]
        if t.empty:
            continue
        algum = True
        titulo = "Venceram" if f == FAIXA_VENCIDO else f"Passaram para {f}"
        partes += [f"<h3>{ICONES[f]} {html.escape(titulo)} ({len(t)})</h3>", _html_tabela(t, COR_FAIXA[f])]
    if not algum:
        partes.append('<p class="vazio">Nenhum contrato mudou de faixa esta semana.</p>')

    if not resumo["vencem_semana"].empty:
        partes += ["<h2>⚠️ Vencem nos próximos 7 dias</h2>", _html_tabela(resumo["vencem_semana"], "A32D2D")]

    partes.append("<h2>Contratos em cada faixa</h2>")
    for f in FAIXAS_EXIBIDAS:
        t = resumo["faixas"][f]
        partes += [f"<h3>{ICONES[f]} {f} ({len(t)})</h3>", _html_tabela(t, COR_FAIXA[f])]

    pf = resumo["por_fiscal"]
    partes.append("<h2>👤 Por fiscal (contratos em alerta)</h2>")
    partes.append(_html_tabela(pf, "444444") if not pf.empty else '<p class="vazio">Nenhum contrato em alerta.</p>')

    return (
        '<!DOCTYPE html><html lang="pt-BR"><head><meta charset="utf-8">'
        f"<title>Resumo semanal de contratos — {hoje:%d/%m/%Y}</title><style>{_CSS}</style></head>"
        f"<body>{''.join(partes)}</body></html>"
    )


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def _estilizar_aba(ws, cor_cabecalho: str, larguras: dict[str, int] | None = None):
    fonte = Font(name="Arial", size=10)
    for linha in ws.iter_rows():
        for cel in linha:
            cel.font = fonte
            cel.alignment = Alignment(vertical="top", wrap_text=True)
    for cel in ws[1]:
        cel.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", start_color=cor_cabecalho)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, cel in enumerate(ws[1], start=1):
        ws.column_dimensions[get_column_letter(i)].width = (larguras or {}).get(cel.value, 16)


def gerar_excel(resumo: dict) -> bytes:
    larguras = {"Tipo": 10, "Nº": 12, "Fornecedor": 38, "Objeto": 55, "Vencimento": 13, "Dias": 7, "Fiscal": 30, "Fiscal ": 30}
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        # --- aba Resumo
        c = resumo["contagens"]
        resumo_df = pd.DataFrame(
            {
                "Faixa": FAIXAS_EXIBIDAS,
                "Contratos": [c[f] for f in FAIXAS_EXIBIDAS],
                "Entraram nos últimos 7 dias": [len(resumo["entraram"][f]) for f in FAIXAS_EXIBIDAS],
            }
        )
        resumo_df.to_excel(writer, sheet_name="Resumo", index=False, startrow=3)
        ws = writer.sheets["Resumo"]
        ws["A1"] = f"Resumo semanal de contratos — posição em {resumo['hoje']:%d/%m/%Y}"
        ws["A1"].font = Font(name="Arial", size=14, bold=True)
        ws["A2"] = f"{resumo['total']} contratos e atas ativos"
        ws["A2"].font = Font(name="Arial", size=10, color="666666")
        for cel in ws[4]:
            cel.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cel.fill = PatternFill("solid", start_color="444444")
            cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for linha in ws.iter_rows(min_row=5, max_row=4 + len(FAIXAS_EXIBIDAS)):
            for cel in linha:
                cel.font = Font(name="Arial", size=10)
                if cel.column > 1:
                    cel.alignment = Alignment(horizontal="center")
            linha[0].fill = PatternFill("solid", start_color=COR_FAIXA[linha[0].value])
            linha[0].font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        ws.row_dimensions[4].height = 30
        for coluna, largura in zip("ABC", (22, 14, 30)):
            ws.column_dimensions[coluna].width = largura

        # --- uma aba por faixa
        for f in FAIXAS_EXIBIDAS:
            resumo["faixas"][f].to_excel(writer, sheet_name=f, index=False)
            _estilizar_aba(writer.sheets[f], COR_FAIXA[f], larguras)

        # --- entraram na semana
        novos = pd.concat(
            [t.assign(Faixa=f) for f, t in resumo["entraram"].items() if not t.empty],
            ignore_index=True,
        ) if any(not t.empty for t in resumo["entraram"].values()) else pd.DataFrame(columns=["Faixa"] + COLUNAS)
        novos = novos[["Faixa"] + COLUNAS]
        novos.to_excel(writer, sheet_name="Entraram na semana", index=False)
        _estilizar_aba(writer.sheets["Entraram na semana"], "444444", {**larguras, "Faixa": 16})

        # --- por fiscal
        pf = resumo["por_fiscal"] if not resumo["por_fiscal"].empty else pd.DataFrame(columns=["Fiscal"] + FAIXAS_EXIBIDAS)
        pf.to_excel(writer, sheet_name="Por fiscal", index=False)
        _estilizar_aba(writer.sheets["Por fiscal"], "444444", {"Fiscal": 34, COLUNA_TOTAL_ALERTA: 16})

    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Salvar em disco (usado pelo script agendado)
# ---------------------------------------------------------------------------
def salvar_arquivos(resumo: dict | None = None, pasta: str | None = None) -> dict[str, str]:
    resumo = resumo or montar_resumo()
    pasta = pasta or os.path.join(PASTA_PROJETO, "resumos")
    os.makedirs(pasta, exist_ok=True)
    base = os.path.join(pasta, f"Resumo_Contratos_{resumo['hoje']:%d-%m-%Y}")

    caminhos = {"xlsx": base + ".xlsx", "html": base + ".html", "txt": base + ".txt"}
    with open(caminhos["xlsx"], "wb") as f:
        f.write(gerar_excel(resumo))
    with open(caminhos["html"], "w", encoding="utf-8") as f:
        f.write(gerar_html(resumo))
    with open(caminhos["txt"], "w", encoding="utf-8") as f:
        f.write(gerar_texto(resumo))
    return caminhos